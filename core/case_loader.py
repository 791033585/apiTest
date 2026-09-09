from pathlib import Path
from typing import Any
import json

from openpyxl import load_workbook

from common.json_utils import parse_json_cell
from core.assertion_models import AssertionRule, ExtractRule
from core.models import ApiCase


class ExcelCaseLoader:
    """Load one worksheet of API cases from an Excel workbook."""

    JSON_FIELDS = {"headers", "params", "data", "json", "files"}
    REQUIRED_FIELDS = {"id", "title", "method", "path"}

    @classmethod
    def load(cls, file_path: str | Path, sheet_name: str | None = None) -> list[ApiCase]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"测试用例文件不存在: {path}")

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError("Excel 测试用例没有表头") from exc

        headers = [str(item).strip().lower() if item is not None else "" for item in raw_headers]
        missing = cls.REQUIRED_FIELDS - set(headers)
        if missing:
            raise ValueError(f"Excel 缺少必要字段: {', '.join(sorted(missing))}")

        cases: list[ApiCase] = []
        for row_number, values in enumerate(rows, start=2):
            row = {
                headers[index]: values[index] if index < len(values) else None
                for index in range(len(headers))
                if headers[index]
            }
            if not any(value is not None and str(value).strip() for value in row.values()):
                continue
            try:
                cases.append(cls._to_case(row))
            except (TypeError, ValueError, KeyError) as exc:
                raise ValueError(f"Excel 第 {row_number} 行解析失败: {exc}") from exc
        workbook.close()
        return cases

    @classmethod
    def _to_case(cls, row: dict[str, Any]) -> ApiCase:
        def text(name: str) -> str:
            value = row.get(name)
            return "" if value is None else str(value).strip()

        def json_field(name: str) -> Any:
            return parse_json_cell(row.get(name), name)

        return ApiCase(
            case_id=text("id"),
            module=text("module"),
            feature=text("feature"),
            story=text("story"),
            title=text("title"),
            method=text("method").upper(),
            path=text("path"),
            headers=json_field("headers"),
            params=json_field("params"),
            data=json_field("data"),
            json_body=json_field("json"),
            files=json_field("files"),
            check=text("check"),
            expected=row.get("expected"),
            extract=text("extract"),
            sql_check=text("sql_check"),
            notes=text("notes"),
            enabled=cls._to_bool(row.get("enabled", True)),
            assertions=cls._assertions(row.get("assertions"), text("check"), row.get("expected")),
            extract_rules=cls._extracts(row.get("extracts"), text("extract")),
            depends_on=cls._string_list(row.get("depends_on")),
            case_type=text("case_type"),
            priority=text("priority"),
        )

    @staticmethod
    def _optional_json(value: Any, field_name: str) -> Any:
        if value is None or str(value).strip() == "":
            return None
        if isinstance(value, (dict, list)):
            return value
        try:
            return json.loads(str(value))
        except json.JSONDecodeError as exc:
            raise ValueError(f"字段 {field_name} 不是合法 JSON") from exc

    @classmethod
    def _assertions(cls, value: Any, legacy_check: str, legacy_expected: Any) -> list[AssertionRule]:
        parsed = cls._optional_json(value, "assertions")
        if parsed is None:
            if not legacy_check:
                return []
            if legacy_check.startswith("[") or legacy_check.startswith("{"):
                parsed = cls._optional_json(legacy_check, "check")
            else:
                return [
                    AssertionRule(
                        target="status_code" if legacy_check.lower() == "status_code" else "json",
                        path=legacy_check,
                        expected=legacy_expected,
                    )
                ]
        if isinstance(parsed, dict):
            parsed = [parsed]
        if not isinstance(parsed, list):
            raise ValueError("字段 assertions 必须是数组或对象")
        return [AssertionRule.from_dict(item) for item in parsed]

    @classmethod
    def _extracts(cls, value: Any, legacy_extract: str) -> list[ExtractRule]:
        parsed = cls._optional_json(value, "extracts")
        if parsed is not None:
            if isinstance(parsed, dict):
                parsed = [parsed]
            if not isinstance(parsed, list):
                raise ValueError("字段 extracts 必须是数组或对象")
            return [ExtractRule.from_dict(item) for item in parsed]
        result = []
        for item in legacy_extract.replace(";", "\n").splitlines():
            item = item.strip()
            if not item:
                continue
            if "=" not in item:
                raise ValueError(f"变量提取格式错误: {item}")
            name, path = [part.strip() for part in item.split("=", 1)]
            result.append(ExtractRule(name=name, path=path))
        return result

    @classmethod
    def _string_list(cls, value: Any) -> list[str]:
        parsed = cls._optional_json(value, "depends_on")
        if parsed is None:
            return []
        if isinstance(parsed, str):
            return [parsed]
        if not isinstance(parsed, list):
            raise ValueError("字段 depends_on 必须是数组")
        return [str(item) for item in parsed]

    @staticmethod
    def _to_bool(value: Any) -> bool:
        if value is None or str(value).strip() == "":
            return True
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() not in {"false", "0", "no", "否"}
